from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Gestión y Normatividad Minera"
    subtitle.text = "Resolución 759 de 2024, Requerimientos FBM y Ley 685 de 2001"

    # Slide 2: Resolución 759 de 2024
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    title_shape.text = "Resolución 759 del 28 de octubre de 2024"

    body_shape = shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Implementación de la Plataforma de Trazabilidad de Minerales (PTM)"

    p = tf.add_paragraph()
    p.text = "Objetivo: Verificación en tiempo real de la procedencia lícita de minerales."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Transparencia: Control de saldos para explotadores y comercializadores."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Obligatoriedad: Registro obligatorio de transacciones vía OTTM."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Integración: Sincronización con AnnA Minería y RUCOM."
    p.level = 1

    # Slide 3: Análisis de Requerimientos Mineros (2024-2025)
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    title_shape.text = "Estadísticas de Requerimientos Mineros"

    body_shape = shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Predominancia de Respuesta a Requerimientos FBM"

    p = tf.add_paragraph()
    p.text = "Año 2024: El FBM representó el 44% de los requerimientos."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Año 2025: El FBM representa el 33% de la participación."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Conclusión: El cumplimiento del Formato Básico Minero es la principal carga administrativa."
    p.level = 1

    # Slide 4: Ley 685 de 2001: Régimen de Transición
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    title_shape.text = "Ley 685 de 2001: Régimen de Transición"

    body_shape = shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Artículos 348 a 351: Seguridad Jurídica"

    p = tf.add_paragraph()
    p.text = "Art. 348: Respeto a la validez de títulos anteriores (D 2655 de 1988)."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Art. 350: Continuidad de condiciones y términos del régimen anterior."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Derechos Adquiridos: Los títulos perfeccionados mantienen sus obligaciones originales."
    p.level = 1

    # Slide 5-7: Inventario de Títulos (Table)
    titles_data = [
        ("5981 (EFLK-01)", "Contrato de Concesión", "L 685"),
        ("T10862011 (FEFO-01)", "Contrato de Concesión", "L 685"),
        ("C3924011 (EAVB-01)", "Contrato de Concesión", "L 685"),
        ("T7763011 (EILE-01)", "Contrato de Concesión", "L 685"),
        ("T11399011 (FFPE-03)", "Contrato de Concesión", "L 685"),
        ("C11414B011 (FFRI-01)", "Por definir", "Por definir"),
        ("C11417011 (FFSA-01)", "Contrato de Concesión", "D 2655"),
        ("11414A ()", "Por definir", "Por definir"),
        ("ICT-16151 (ICT-16151)", "Contrato de Concesión", "L 685"),
        ("H766B005 (CFWB-02)", "Contrato de Concesión", "L 685"),
        ("T663005 (GCNM-06)", "Licencia de Explotación", "D 2655"),
        ("4405 (ECGD-01)", "Contrato de Concesión", "L 685"),
        ("T4925005 (HCIJ-10)", "Contrato de Concesión", "L 685"),
        ("C4452011 (ECKH-01)", "Contrato de Concesión", "L 685"),
        ("FLD-157 (FLD-157)", "Contrato de Concesión", "L 685"),
        ("JAE-11231 (JAE-11231)", "Contrato de Concesión", "L 685"),
        ("00033-15 ()", "Contrato de Concesión", "L 685"),
        ("8960 (FAIN-01)", "Contrato de Concesión", "L 685"),
        ("14585 (GAL-01)", "Contrato de Concesión", "L 685"),
        ("15929 (GCAO-04)", "Contrato de Concesión", "L 685"),
        ("00141-15 (GFXE-04)", "Licencia de Explotación", "D 2655"),
        ("LJ7-09361 (LJ7-09361)", "Contrato de Concesión", "L 685"),
        ("00918-15 (ADD-161)", "Contrato de Concesión", "L 685"),
        ("00988-15 (HFVJ-08)", "Contrato de Concesión", "L 685"),
        ("IJ5-08181 (IJ5-08181)", "Contrato de Concesión", "L 685"),
        ("3450 (DJLF-01)", "Contrato de Concesión", "L 685"),
        ("13549 (FIOI-02)", "Contrato de Concesión", "L 685"),
        ("22212 (GIKC-01)", "Contrato de Concesión", "L 685"),
        ("3871 (EARJ-01)", "Contrato de Concesión", "L 685"),
        ("9098 (FANO-01)", "Contrato de Concesión", "L 685"),
        ("EJ3-101 (EJ3-101)", "Contrato de Concesión", "D 2655"),
        ("13101 (FJKB-01)", "Contrato de Concesión", "L 685"),
        ("1995 (DEDJ-01)", "Contrato de Concesión", "L 685"),
        ("1871 (DDTK-01)", "Contrato de Concesión", "D 2655"),
    ]

    rows_per_slide = 12
    for i in range(0, len(titles_data), rows_per_slide):
        slide = prs.slides.add_slide(prs.slide_layouts[5]) # Blank layout with title
        slide.shapes.title.text = f"Inventario de Títulos Mineros (Parte {i//rows_per_slide + 1})"

        subset = titles_data[i:i+rows_per_slide]
        rows = len(subset) + 1
        cols = 3
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9.0)
        height = Inches(0.8)

        table = slide.shapes.add_table(rows, cols, left, top, width, height).table

        # Headers
        table.cell(0, 0).text = "Placa Título"
        table.cell(0, 1).text = "Tipo de Título"
        table.cell(0, 2).text = "Régimen"

        for row_idx, (placa, tipo, regimen) in enumerate(subset):
            table.cell(row_idx + 1, 0).text = placa
            table.cell(row_idx + 1, 1).text = tipo
            table.cell(row_idx + 1, 2).text = regimen

            # Formatting
            for col_idx in range(3):
                cell = table.cell(row_idx + 1, col_idx)
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

    prs.save("presentacion_minera.pptx")
    print("Presentación generada exitosamente: presentacion_minera.pptx")

if __name__ == "__main__":
    create_presentation()
